import argparse
import calendar
import os
import re
from dataclasses import dataclass
from datetime import datetime
from typing import Any

import pandas as pd
import stripe
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from prettytable import PrettyTable

from cache import StripeCache

stripe.api_version = "2023-10-16"


# ── Stripe client with transparent caching ──────────────────────────────────────

class StripeClient:
    def __init__(self, cache: StripeCache, debug: bool = False):
        self._cache = cache
        self._debug = debug

    def _fetch(self, resource_cls, id: str) -> Any:
        obj = self._cache.get(id)
        if obj is None:
            obj = resource_cls.retrieve(id)
            self._cache.set(id, obj)
        elif self._debug:
            print(f"  cache hit: {id}")
        if "id" not in obj or obj.id != id:
            raise ValueError(f"{resource_cls.__name__} ID mismatch: got {obj.id!r}, expected {id!r}")
        return obj

    def product(self, id: str):         return self._fetch(stripe.Product,       id)
    def invoice(self, id: str):         return self._fetch(stripe.Invoice,        id)
    def payment_intent(self, id: str):  return self._fetch(stripe.PaymentIntent,  id)
    def charge(self, id: str):          return self._fetch(stripe.Charge,         id)
    def refund(self, id: str):          return self._fetch(stripe.Refund,         id)
    def dispute(self, id: str):         return self._fetch(stripe.Dispute,        id)


# ── Pagination ──────────────────────────────────────────────────────────────────

def paginate(list_fn, **kwargs) -> list:
    """Exhaust all pages from a Stripe list endpoint and return every item."""
    page = list_fn(limit=100, **kwargs)
    items = list(page.data)
    while page.has_more:
        page = list_fn(limit=100, starting_after=page.data[-1].id, **kwargs)
        items.extend(page.data)
    return items


# ── Data structures ─────────────────────────────────────────────────────────────

@dataclass
class TransactionRecord:
    type: str            # "charge" | "refund" | "dispute" | "dispute_reversal"
    product: str
    customer_email: str
    customer_name: str
    amount: int         # cents
    fee: int            # cents
    net: int            # cents
    payment_intent: str
    receipt_url: str
    transaction_id: str


class RevenueAccumulator:
    """Tracks per-product revenue (integer cents) and detailed transaction records."""

    def __init__(self):
        self._products: dict[str, Any] = {}
        self._revenue: dict[str, int] = {}
        self.transaction_records: dict[str, list[TransactionRecord]] = {}

    def add(self, product: Any, net_cents: int) -> None:
        assert isinstance(net_cents, int), (
            f"Revenue must be tracked as integer cents, got {type(net_cents).__name__}: {net_cents!r}"
        )
        pid = product["id"]
        if pid not in self._revenue:
            self._products[pid] = product
            self._revenue[pid] = 0
        self._revenue[pid] += net_cents

    def add_system(self, label: str, net_cents: int) -> None:
        """Add a non-product line (e.g. fee, payout_minimum_balance_*)."""
        self.add({"id": label, "name": label}, net_cents)

    def add_transaction_record(self, product_name: str, record: TransactionRecord) -> None:
        self.transaction_records.setdefault(product_name, []).append(record)

    def entries(self) -> list[tuple[Any, int]]:
        return [(self._products[pid], rev) for pid, rev in self._revenue.items()]

    def total_cents(self) -> int:
        return sum(self._revenue.values())


# ── Product resolution ──────────────────────────────────────────────────────────

def _checkout_product(client: StripeClient, session_id: str) -> Any:
    """Return the product from a single-item checkout session."""
    items = stripe.checkout.Session.list_line_items(session_id)
    assert len(items.data) == 1, f"Expected 1 line item in session {session_id}, got {len(items.data)}"
    return client.product(items.data[0].price.product)


def _invoice_product(client: StripeClient, invoice_id: str) -> tuple[Any, Any]:
    """Return (product, invoice) for a single-line invoice."""
    invoice = client.invoice(invoice_id)
    assert len(invoice.lines.data) == 1, f"Expected 1 invoice line, got {len(invoice.lines.data)}"
    return client.product(invoice.lines.data[0].price.product), invoice


def resolve_from_charge(client: StripeClient, charge) -> tuple[Any, str, str]:
    """
    Return (product, customer_email, customer_name) for a charge.
    Supports subscription (invoice) and one-time (checkout session) flows.
    """
    if getattr(charge, "invoice", None):
        product, invoice = _invoice_product(client, charge.invoice)
        return product, invoice.customer_email or "", invoice.customer_name or ""

    pi = client.payment_intent(charge.payment_intent)
    sessions = stripe.checkout.Session.list(payment_intent=pi.id)
    if sessions.data:
        session = sessions.data[0]
        product = _checkout_product(client, session.id)
        cd = session.customer_details
        return product, (cd.email or "" if cd else ""), (cd.name or "" if cd else "")

    raise ValueError(f"Cannot resolve product for charge {charge.id}: no invoice or checkout session found")


def resolve_from_refund(client: StripeClient, refund) -> tuple[Any, str, str, str, str]:
    """Return (product, customer_email, customer_name, receipt_url, payment_intent) for a refund."""
    pi = client.payment_intent(refund.payment_intent)

    if getattr(pi, "payment_details", None) and pi.payment_details.order_reference.startswith("cs_"):
        session = stripe.checkout.Session.retrieve(pi.payment_details.order_reference)
        product = _checkout_product(client, session.id)
        cd = session.customer_details
        email, name = (cd.email or "" if cd else ""), (cd.name or "" if cd else "")
    elif getattr(pi, "invoice", None):
        product, invoice = _invoice_product(client, pi.invoice)
        email, name = invoice.customer_email or "", invoice.customer_name or ""
    else:
        raise ValueError(f"Cannot resolve product for refund {refund.id}")

    charge = client.charge(refund.charge) if getattr(refund, "charge", None) else None
    receipt_url = (getattr(charge, "receipt_url", "") or "") if charge else ""
    return product, email, name, receipt_url, refund.payment_intent or ""


def resolve_from_dispute(client: StripeClient, dispute) -> tuple[Any, str, str, str, str]:
    """Return (product, customer_email, customer_name, receipt_url, payment_intent) for a dispute or dispute_reversal."""
    charge = client.charge(dispute.charge)
    product, email, name = resolve_from_charge(client, charge)
    receipt_url = getattr(charge, "receipt_url", "") or ""
    return product, email, name, receipt_url, charge.payment_intent or ""


# ── Transaction processing ──────────────────────────────────────────────────────

def process_transaction(t, client: StripeClient, acc: RevenueAccumulator) -> None:
    """
    Attribute one non-payout balance transaction's net to the correct product bucket.
    Raises on any unrecognised type with a non-zero net so the books always balance.
    """
    # Stripe guarantees net == amount - fee for every balance transaction.
    # If this fails the source data itself is corrupt and nothing downstream is trustworthy.
    assert t.net == t.amount - t.fee, (
        f"Transaction {t.id} ({t.reporting_category}): "
        f"net {t.net} != amount {t.amount} - fee {t.fee} = {t.amount - t.fee}"
    )

    cat = t.reporting_category

    if cat.startswith("payout_minimum_balance_"):
        acc.add_system(cat, t.net)

    elif cat == "fee":
        acc.add_system("fee", t.net)

    elif cat == "refund":
        assert t.source, f"refund transaction {t.id} has no source"
        product, email, name, receipt_url, pi_id = resolve_from_refund(client, client.refund(t.source))
        acc.add(product, t.net)
        acc.add_transaction_record(product["name"], TransactionRecord(
            type="refund",
            product=product["name"],
            customer_email=email,
            customer_name=name,
            amount=t.amount,
            fee=t.fee,
            net=t.net,
            payment_intent=pi_id,
            receipt_url=receipt_url,
            transaction_id=t.id,
        ))

    elif cat == "charge":
        assert t.source, f"charge transaction {t.id} has no source"
        charge = client.charge(t.source)
        assert charge.payment_intent, f"charge {charge.id} has no payment_intent"
        product, email, name = resolve_from_charge(client, charge)
        acc.add(product, t.net)
        acc.add_transaction_record(product["name"], TransactionRecord(
            type="charge",
            product=product["name"],
            customer_email=email,
            customer_name=name,
            amount=t.amount,
            fee=t.fee,
            net=t.net,
            payment_intent=charge.payment_intent,
            receipt_url=getattr(charge, "receipt_url", "") or "",
            transaction_id=t.id,
        ))

    elif cat in ("dispute", "dispute_reversal"):
        # dispute: money debited when a dispute is opened against you
        # dispute_reversal: money returned when a dispute is won or closed
        assert t.source, f"{cat} transaction {t.id} has no source"
        product, email, name, receipt_url, pi_id = resolve_from_dispute(client, client.dispute(t.source))
        acc.add(product, t.net)
        acc.add_transaction_record(product["name"], TransactionRecord(
            type=cat,
            product=product["name"],
            customer_email=email,
            customer_name=name,
            amount=t.amount,
            fee=t.fee,
            net=t.net,
            payment_intent=pi_id,
            receipt_url=receipt_url,
            transaction_id=t.id,
        ))

    elif t.net != 0:
        raise RuntimeError(
            f'Unhandled transaction type "{cat}" with non-zero net {t.net} for {t.id}. '
            "Add explicit handling to keep all revenue attributed to a product."
        )


# ── Reporting ───────────────────────────────────────────────────────────────────

def build_table(acc: RevenueAccumulator, account: str, year: int, month: int) -> PrettyTable:
    fee_cents = next((rev for prod, rev in acc.entries() if prod["name"] == "fee"), 0)
    payout_min_bal_cents = sum(
        rev for prod, rev in acc.entries() if prod["name"].startswith("payout_minimum_balance_")
    )
    total_product_cents = sum(
        rev for prod, rev in acc.entries()
        if not prod["name"].startswith("payout_minimum_balance_") and prod["name"] != "fee"
    )

    # payout_minimum_balance holds/releases must net to zero — they are internal
    # Stripe mechanics, not real revenue, and are excluded from the table entirely.
    assert payout_min_bal_cents == 0, (
        f"payout_minimum_balance entries do not net to zero: {payout_min_bal_cents} cents. "
        f"Investigate before trusting any revenue figures."
    )
    # With min_bal == 0, the three buckets simplify: fee + products == total payout.
    assert fee_cents + total_product_cents == acc.total_cents(), (
        f"fee({fee_cents}) + products({total_product_cents}) = {fee_cents + total_product_cents} "
        f"!= total payout {acc.total_cents()}"
    )

    table = PrettyTable(["Product", "Revenue ($)", "Adjusted Fee ($)", "Adjusted Revenue ($)", "Email", "Rate (%)", "Net Payout ($)"])
    table.title = (
        f"Stripe Revenue Report for {account} {year}-{month:02d}"
        " (generated by https://github.com/batchfy/stripe-revenue-report)"
    )

    total_adj_rev = 0.0
    total_adj_fee = 0.0
    total_revenue_sum = 0   # integer cents: sum of every Revenue($) cell displayed

    for prod, rev in acc.entries():
        name = prod["name"]
        if name.startswith("payout_minimum_balance_"):
            continue

        if name == "fee":
            adj_fee, adj_rev = 0.0, 0.0
        else:
            adj_fee = fee_cents * rev / total_product_cents if total_product_cents else 0.0
            adj_rev = rev + adj_fee

        total_adj_rev += adj_rev
        total_adj_fee += adj_fee
        total_revenue_sum += rev

        metadata = prod["metadata"] if "metadata" in prod else {}
        email = metadata["email"] if "email" in metadata else ""
        rate_raw = metadata["rate"] if "rate" in metadata else None
        try:
            rate: float | None = float(rate_raw) if rate_raw is not None else None
        except ValueError:
            rate = None

        net_payout = adj_rev * (1 - rate) if rate is not None else adj_rev
        rate_display = f"{rate * 100:.1f}" if rate is not None else ""

        table.add_row([
            name,
            f"{rev / 100:.2f}",
            f"{adj_fee / 100:.2f}",
            f"{adj_rev / 100:.2f}",
            email,
            rate_display,
            f"{net_payout / 100:.2f}",
        ])

    # ── Check: Revenue($) column — all displayed rows sum exactly to total payout ─
    # Valid because payout_min_bal == 0 was asserted above.
    assert total_revenue_sum == acc.total_cents(), (
        f"Sum of Revenue($) column {total_revenue_sum} != total payout {acc.total_cents()} "
        f"(gap: {acc.total_cents() - total_revenue_sum})"
    )
    # ── Check: fee was fully distributed across products ──────────────────────────
    assert abs(total_adj_fee - fee_cents) < 1, (
        f"Fee distribution error: distributed {total_adj_fee:.4f} != fee total {fee_cents} "
        f"(delta: {total_adj_fee - fee_cents:.6f})"
    )
    # ── Check: Adjusted Revenue total == total payout (fee absorbed into products) ─
    assert abs(total_adj_rev - acc.total_cents()) < 1, (
        f"Adjusted Revenue total {total_adj_rev:.4f} != total payout {acc.total_cents()} "
        f"(delta: {total_adj_rev - acc.total_cents():.6f})"
    )

    table.add_row(["Total",
                   f"{total_revenue_sum / 100:.2f}",
                   f"{total_adj_fee / 100:.2f}",
                   f"{total_adj_rev / 100:.2f}",
                   "", "", ""])
    return table


_UNSAFE_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\[\]]')


def _safe_filename(name: str) -> str:
    return _UNSAFE_FILENAME_CHARS.sub("_", name).strip()


def _mask_email(email: str) -> str:
    if not email or "@" not in email:
        return email
    local, domain = email.split("@", 1)
    masked_local = f"{local[0]}***{local[-1]}" if len(local) > 1 else f"{local}***"
    return f"{masked_local}@{domain}"


def _records_to_df(records: list[TransactionRecord]) -> pd.DataFrame:
    return pd.DataFrame([
        {"type": r.type, "product": r.product, "customer_email": _mask_email(r.customer_email),
         "customer_name": r.customer_name,
         "amount": r.amount, "fee": r.fee, "net": r.net,
         "payment_intent": r.payment_intent, "receipt_url": r.receipt_url, "transaction_id": r.transaction_id}
        for r in records
    ])


_HYPERLINK_FONT = Font(color="0563C1", underline="single")


def _linkify_column(worksheet, df: pd.DataFrame, column: str) -> None:
    """Turn a column of plain URL strings into clickable Excel hyperlinks."""
    if column not in df.columns:
        return
    col_idx = df.columns.get_loc(column) + 1  # openpyxl columns are 1-indexed
    for row_idx, url in enumerate(df[column], start=2):  # row 1 is the header
        if not url:
            continue
        cell = worksheet.cell(row=row_idx, column=col_idx)
        cell.hyperlink = url
        cell.font = _HYPERLINK_FONT


_COLUMN_WIDTHS = {
    "type": 12,
    "product": 20,
    "customer_email": 22,
    "customer_name": 20,
    "amount": 10,
    "fee": 10,
    "net": 10,
    "payment_intent": 20,
    "receipt_url": 60,
    "transaction_id": 20,
}
_DEFAULT_COLUMN_WIDTH = 15


def _set_column_widths(worksheet, df: pd.DataFrame) -> None:
    for col_idx, col_name in enumerate(df.columns, start=1):
        width = _COLUMN_WIDTHS.get(col_name, _DEFAULT_COLUMN_WIDTH)
        worksheet.column_dimensions[get_column_letter(col_idx)].width = width


def save_outputs(acc: RevenueAccumulator, table: PrettyTable, account: str, year: int, month: int) -> None:
    month_str = f"{year}-{month:02d}"
    out_dir = os.path.join("reports", _safe_filename(account), month_str)
    os.makedirs(out_dir, exist_ok=True)

    for prod_name, records in acc.transaction_records.items():
        safe_name = _safe_filename(prod_name)
        sheet_name = safe_name[:31]
        path = os.path.join(out_dir, f"{month_str}_{safe_name}.xlsx")
        with pd.ExcelWriter(path) as w:
            df = _records_to_df(records)
            df.to_excel(w, sheet_name=sheet_name, index=False)
            _linkify_column(w.sheets[sheet_name], df, "receipt_url")
            _set_column_widths(w.sheets[sheet_name], df)


# ── Utilities ───────────────────────────────────────────────────────────────────

def month_range(year: int, month: int) -> tuple[datetime, datetime]:
    start = datetime(year, month, 1)
    end = datetime(year + 1, 1, 1) if month == 12 else datetime(year, month + 1, 1)
    return start, end


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Calculate Stripe revenue for a specific month.")
    parser.add_argument("--year",  type=int, default=datetime.now().year)
    parser.add_argument("--month", type=int, default=datetime.now().month)
    parser.add_argument("--debug", action="store_true")
    parser.add_argument("--payout-only", action="store_true", help="Only count this month's payouts and exit")
    parser.add_argument("--delete-cache", metavar="KEY", help="Delete a cache entry by key and exit")
    return parser.parse_args()


def account_name() -> str:
    """Derive a human-readable account name from the API key's Stripe account."""
    account = stripe.Account.retrieve()
    settings = account["settings"] if "settings" in account else None
    dashboard = settings["dashboard"] if settings and "dashboard" in settings else None
    display_name = dashboard["display_name"] if dashboard and "display_name" in dashboard else None

    business_profile = account["business_profile"] if "business_profile" in account else None
    business_name = business_profile["name"] if business_profile and "name" in business_profile else None

    email = account["email"] if "email" in account else None

    return display_name or business_name or email or account.id


# ── Entrypoint ──────────────────────────────────────────────────────────────────

def main() -> None:
    args = parse_args()
    stripe.api_key = os.getenv("STRIPE_SECRET_KEY")

    with StripeCache() as cache:
        client = StripeClient(cache, debug=args.debug)

        if args.delete_cache:
            print("Deleted." if cache.delete(args.delete_cache) else "Key not found.")
            return

        account = account_name()
        print(f"Account: {account}")

        start, end = month_range(args.year, args.month)
        print(f"Fetching payouts {start:%Y-%m-%d} → {end:%Y-%m-%d} …")

        payouts = paginate(
            stripe.Payout.list,
            created={"gte": int(start.timestamp()), "lt": int(end.timestamp())},
        )
        total_amount = sum(p.amount for p in payouts)
        print(f"{len(payouts)} payout(s), total {total_amount / 100:.2f}")

        if args.payout_only:
            return

        # Fetch all transactions per payout and verify the raw sums immediately.
        payout_transactions: dict[str, list] = {}
        for payout in payouts:
            ts = paginate(stripe.BalanceTransaction.list, payout=payout.id)
            non_payout_net = sum(t.net for t in ts if t.type != "payout")
            assert non_payout_net == payout.amount, (
                f"Payout {payout.id}: transaction net sum {non_payout_net} "
                f"!= payout amount {payout.amount}"
            )
            payout_transactions[payout.id] = ts

        # Attribute every transaction to a product bucket.
        acc = RevenueAccumulator()
        accumulated = {po.id: 0 for po in payouts}
        running_total = 0  # cumulative payout amounts fully processed so far

        for payout in payouts:
            for t in payout_transactions[payout.id]:
                if t.type == "payout":
                    continue
                process_transaction(t, client, acc)
                accumulated[payout.id] += t.net
                print(
                    f"  [{payout.id}] {t.id}  {t.reporting_category:<30}"
                    f"  net={t.net:+9d}  running={accumulated[payout.id]}/{payout.amount}"
                )

            # ── Per-payout closure check (runs as soon as each payout finishes) ──
            assert accumulated[payout.id] == payout.amount, (
                f"Payout {payout.id}: attributed net {accumulated[payout.id]} "
                f"!= payout amount {payout.amount} "
                f"(gap: {payout.amount - accumulated[payout.id]})"
            )
            # ── Progressive global check (acc grows payout-by-payout) ────────────
            running_total += payout.amount
            assert acc.total_cents() == running_total, (
                f"After processing payout {payout.id}: "
                f"acc.total_cents() {acc.total_cents()} != running total {running_total} "
                f"(gap: {running_total - acc.total_cents()})"
            )
            print(f"  ✓ Payout {payout.id}: {payout.amount / 100:.2f} reconciles exactly.")

        # ── Final belt-and-suspenders check across all payouts ───────────────────
        assert acc.total_cents() == total_amount, (
            f"Revenue total {acc.total_cents()} != sum of all payouts {total_amount} "
            f"(gap: {total_amount - acc.total_cents()})"
        )
        print(f"Balance verified: all {len(payouts)} payout(s) reconcile exactly.")

        table = build_table(acc, account, args.year, args.month)
        print(table)
        save_outputs(acc, table, account, args.year, args.month)


if __name__ == "__main__":
    main()
